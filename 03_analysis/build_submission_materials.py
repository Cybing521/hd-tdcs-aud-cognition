#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build cover letter and submission helper documents."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "02_投稿材料"
SENS_PATH = ROOT / "04_data" / "processed" / "sensitivity_analyses.json"
TITLE_EN = (
    "Effects of High-Definition Transcranial Direct Current Stimulation on Cognitive "
    "Function in Male Inpatients with Alcohol Dependence: A Randomized Sham-Controlled Trial"
)


def write_cover_letter() -> None:
    mde = 0.83
    if SENS_PATH.exists():
        mde = json.loads(SENS_PATH.read_text(encoding="utf-8")).get(
            "minimum_detectable_effect", {}
        ).get("min_detectable_d_80power_21v27", 0.83)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p = doc.add_paragraph()
    p.add_run("Cover Letter — Brain Stimulation").bold = True
    doc.add_paragraph("")
    doc.add_paragraph("Dear Editor,")
    doc.add_paragraph(
        f'We submit the original research article entitled "{TITLE_EN}" for consideration '
        "as an Original Research article in Brain Stimulation."
    )
    doc.add_paragraph(
        "Why this journal. The study tests a focal neuromodulation montage—bilateral 4×1 "
        "HD-tDCS over the DLPFC—in a sham-controlled three-arm inpatient trial, with a "
        "prespecified standardized cognitive battery rather than craving or relapse alone. "
        "That combination sits at the interface of stimulation technology, prefrontal targeting, "
        "and cognitive outcome measurement that Brain Stimulation routinely publishes."
    )
    doc.add_paragraph(
        "Key findings. Active HD-tDCS produced medium-to-large advantages versus sham on "
        "TMT-B, VFT, and HVLT-R (|Cohen's d| = 0.96–1.30; Hedges g = −0.94/1.03/1.28). "
        "Baseline-adjusted ANCOVA confirmed all three primary outcomes (partial η² = 0.19–0.31). "
        "Sensitivity analyses that altered assignment of ambiguous post-randomization records "
        f"preserved these conclusions (ANCOVA P≤0.0003). The minimum detectable d at 80% power "
        f"for n=21 versus n=27 is {mde:.2f}; observed effects exceeded that threshold. Sham did "
        "not differ from routine care."
    )
    doc.add_paragraph(
        "Reporting. The manuscript includes a CONSORT flow with explicit reassignment accounting "
        "(Figure 1), an HD-tDCS montage schematic (Figure 2), outcome and trajectory figures "
        "(Figures 3–4), a structured abstract, Integrity of the Blind, FDR-controlled primary "
        "omnibus tests, effect sizes with 95% confidence intervals and Hedges g, and transparent "
        "disclosure of data verification and the absence of structured adverse-event collection "
        "in the original protocol."
    )
    doc.add_paragraph(
        "Trial registration. The study was not prospectively registered before enrollment (2019); "
        "this limitation is disclosed, and retrospective registration with disclosure of the "
        "registration date is planned before final upload. No overlapping publication is planned."
    )
    doc.add_paragraph(
        "All authors approve this submission. Author names, affiliations, funding, ethics approval "
        "numbers, and conflict-of-interest statements will be completed in the final submission package."
    )
    doc.add_paragraph("")
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph("[Corresponding author name and signature]")
    doc.add_paragraph("[Institution]")
    doc.add_paragraph("[Date]")

    out = OUT_DIR / "Cover_letter_Brain_Stimulation.docx"
    doc.save(out)
    print("Wrote", out)


def write_highlights_txt() -> None:
    from build_full_manuscript import HIGHLIGHTS_EN

    lines = ["Highlights (Brain Stimulation submission form)", ""]
    lines += [f"• {b}" for b in HIGHLIGHTS_EN]
    text = "\n".join(lines) + "\n"
    out = OUT_DIR / "Highlights.txt"
    out.write_text(text, encoding="utf-8")
    print("Wrote", out)


def write_submission_checklist() -> None:
    text = """# Submission Package Checklist — Brain Stimulation (locked)

## Manuscript files
- [ ] Main manuscript (double-spaced BRS layout): `01_论文正文/HD-tDCS_alcohol_dependence_cognition_BRS.docx`
- [ ] English working copy: `01_论文正文/HD-tDCS_alcohol_dependence_cognition_EN.docx`
- [ ] Chinese working manuscript: `01_论文正文/HD-tDCS酒依赖认知功能_修订稿.docx`

## Figures (upload PDF preferred; PNG backup)
- [ ] Figure 1 CONSORT: `01_论文正文/figures/Figure1_CONSORT_flowchart.pdf`
- [ ] Figure 2 Montage: `01_论文正文/figures/Figure2_HD_tDCS_montage.pdf`
- [ ] Figure 3 Post outcomes: `01_论文正文/figures/Figure3_primary_outcomes_post.pdf`
- [ ] Figure 4 Change trajectories: `01_论文正文/figures/Figure4_change_trajectories.pdf`

## Editorial Manager extras
- [ ] Cover letter: `02_投稿材料/Cover_letter_Brain_Stimulation.docx`
- [ ] Highlights (≤85 chars each): `02_投稿材料/Highlights.txt`
- [ ] Keywords (in manuscript)
- [ ] Financial disclosures / COI
- [ ] Generative AI declaration (in manuscript, above References)
- [ ] Integrity of the Blind section (in Methods)

## QC
- [ ] Primary review: `ccfa-review-reports/manuscript_review_v3_2026-08-24.md`
- [ ] CONSORT checklist: `ccfa-review-reports/consort_checklist.md`
- [ ] Reference validation: `ccfa-review-reports/ref_validation_2026-08-24.md`
- [ ] Workflow state: `ccfa.yaml`

## Author TODO before upload
- [ ] Author list, degrees, affiliations, corresponding author contact
- [ ] Ethics approval number
- [ ] Funding + CRediT
- [ ] Retrospective trial registration number + registration date
- [ ] Allocation concealment detail; electrode size/material; AE chart review if available
"""
    out = OUT_DIR / "Submission_checklist.md"
    out.write_text(text, encoding="utf-8")
    print("Wrote", out)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_cover_letter()
    write_highlights_txt()
    write_submission_checklist()


if __name__ == "__main__":
    main()
