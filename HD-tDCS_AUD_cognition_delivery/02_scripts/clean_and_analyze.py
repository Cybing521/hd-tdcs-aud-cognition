#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Clean duplicate IDs, reassign sham-group cases, recompute tables."""

from __future__ import annotations

import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
from scipy import stats
from statsmodels.stats.multicomp import pairwise_tukeyhsd
from statsmodels.stats.oneway import anova_oneway

ROOT = Path(__file__).resolve().parents[1]
TABLE_DIR = ROOT / "03_分析表格"
OUT_DIR = ROOT / "04_data" / "processed"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Cross-group mis-assigned cases: keep sham only.
SHAM_ONLY = {"P-A", "P-B", "P-C", "P-D", "P-E"}

GROUP_LABEL = {1: "active", 2: "sham", 3: "control"}
GROUP_CN = {1: "正常刺激组", 2: "伪刺激组", 3: "对照组"}

DEMO_COLS = ["年龄", "教育年限", "首次饮酒年龄", "累计使用时间", "成瘾时间", "戒酒次数"]
COG_BASE = {
    "MMSE": ("MMSE1", "MMSE2", "MMSE差值"),
    "TMT-B": ("TMT-B1", "TMT-B2", "TMT-B差值"),
    "DST": ("DST1", "DST2", "DST差值"),
    "VFT": ("VFT1", "VFT2", "VFT差值"),
    "HVLT-R": ("HVLT-R1", "HVLT-R2", "HVLT-R差值"),
}


def load_compare() -> tuple[list[str], list[dict[str, Any]]]:
    path = TABLE_DIR / "比较.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(headers, row)))
    wb.close()
    return headers, rows


def save_compare(headers: list[str], rows: list[dict[str, Any]]) -> None:
    path = TABLE_DIR / "比较.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    wb.save(path)


def pick_lei_ang(rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    cand = [r for r in rows if r.get("姓名") == "P-E"]
    if not cand:
        return None
    # Prefer age 38 row; fallback first.
    cand.sort(key=lambda x: (x.get("年龄") != 38, x.get("年龄") or 999))
    return dict(cand[0])


def clean_compare_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict]:
    audit = {"removed": [], "added": [], "reassigned": []}
    cleaned: list[dict[str, Any]] = []
    lei = pick_lei_ang(rows)

    for r in rows:
        name = r.get("姓名")
        grp = r.get("分组")
        if name in SHAM_ONLY and grp == 1:
            audit["removed"].append({"姓名": name, "原分组": grp, "原因": "应为伪刺激组，删除真刺激重复行"})
            continue
        cleaned.append(r)

    # Ensure single P-E in sham
    cleaned = [r for r in cleaned if not (r.get("姓名") == "P-E")]
    if lei is not None:
        lei["分组"] = 2
        cleaned.append(lei)
        audit["reassigned"].append({"姓名": "P-E", "新分组": 2, "保留年龄": lei.get("年龄")})

    cleaned.sort(key=lambda x: (x.get("分组") or 9, str(x.get("姓名"))))
    return cleaned, audit


def mean_sd(arr: list[float]) -> tuple[float, float, int]:
    vals = [float(x) for x in arr if x is not None and x != "" and not (isinstance(x, float) and math.isnan(x))]
    if not vals:
        return float("nan"), float("nan"), 0
    m = float(np.mean(vals))
    sd = float(np.std(vals, ddof=1)) if len(vals) > 1 else 0.0
    return m, sd, len(vals)


def fmt_ms(m: float, sd: float) -> str:
    if math.isnan(m):
        return "NA"
    return f"{m:.2f}±{sd:.2f}"


def subset(rows: list[dict], grp: int) -> list[dict]:
    return [r for r in rows if r.get("分组") == grp]


def oneway(values_by_group: dict[int, list[float]]) -> dict:
    groups = []
    for k in sorted(values_by_group):
        vals = [float(x) for x in values_by_group[k] if x is not None and x != ""]
        groups.append(vals)
    if any(len(g) < 2 for g in groups):
        return {"F": None, "P": None, "method": "insufficient n", "equal_var": True}
    f_stat, p_val = stats.f_oneway(*groups)
    # Levene for equal variance
    lev = stats.levene(*groups)
    equal_var = lev.pvalue >= 0.05
    return {
        "F": float(f_stat),
        "P": float(p_val),
        "levene_p": float(lev.pvalue),
        "equal_var": equal_var,
        "method": "ANOVA" if equal_var else "Welch",
    }


def welch_anova(values_by_group: dict[int, list[float]]) -> float:
    try:
        res = anova_oneway([values_by_group[k] for k in sorted(values_by_group)], use_var="unequal")
        return float(res.pvalue)
    except Exception:
        return float("nan")


def pairwise(values: list[float], labels: list[int], equal_var: bool) -> dict[tuple[int, int], float]:
    out: dict[tuple[int, int], float] = {}
    pairs = [(1, 2), (1, 3), (2, 3)]
    for a, b in pairs:
        va = [v for v, g in zip(values, labels) if g == a]
        vb = [v for v, g in zip(values, labels) if g == b]
        if len(va) < 2 or len(vb) < 2:
            out[(a, b)] = float("nan")
            continue
        if equal_var:
            _, p = stats.ttest_ind(va, vb, equal_var=True)
        else:
            _, p = stats.ttest_ind(va, vb, equal_var=False)
        out[(a, b)] = float(p)
    return out


def fdr_bh(p_by_outcome: dict[str, float]) -> dict[str, float]:
    """Benjamini-Hochberg FDR adjustment for post-intervention omnibus P values."""
    items = sorted(p_by_outcome.items(), key=lambda x: x[1])
    m = len(items)
    if m == 0:
        return {}
    adjusted: dict[str, float] = {}
    prev = 1.0
    for i in range(m - 1, -1, -1):
        outcome, p_val = items[i]
        val = min(prev, p_val * m / (i + 1))
        adjusted[outcome] = round(float(val), 4)
        prev = val
    return adjusted


def cohen_d(a: list[float], b: list[float]) -> float:
    a = np.array(a, dtype=float)
    b = np.array(b, dtype=float)
    if len(a) < 2 or len(b) < 2:
        return float("nan")
    na, nb = len(a), len(b)
    va, vb = np.var(a, ddof=1), np.var(b, ddof=1)
    pooled = math.sqrt(((na - 1) * va + (nb - 1) * vb) / (na + nb - 2))
    if pooled == 0:
        return 0.0
    return float((np.mean(a) - np.mean(b)) / pooled)


def build_tables(rows: list[dict]) -> dict:
    n = {g: len(subset(rows, g)) for g in (1, 2, 3)}
    result: dict[str, Any] = {"n": n, "tables": {}, "pairwise_post": {}, "correlation": {}, "effect_sizes": {}}

    # Table 1 demographics
    t1 = []
    for col in DEMO_COLS:
        by = {g: [r.get(col) for r in subset(rows, g)] for g in (1, 2, 3)}
        ow = oneway(by)
        if ow["method"] == "insufficient n":
            p = None
        elif ow["equal_var"]:
            p = ow["P"]
        else:
            p = welch_anova(by)
        m1, s1, _ = mean_sd(by[1])
        m2, s2, _ = mean_sd(by[2])
        m3, s3, _ = mean_sd(by[3])
        row = {
            "variable": col,
            "active": fmt_ms(m1, s1),
            "sham": fmt_ms(m2, s2),
            "control": fmt_ms(m3, s3),
            "F": round(ow["F"], 3) if ow.get("F") is not None else None,
            "P": round(p, 3) if p is not None and p == p else None,
            "method": ow["method"],
        }
        t1.append(row)
    result["tables"]["table1"] = t1

    # Baseline & post cognitive
    for phase, col_suffix in [("baseline", "1"), ("post", "2")]:
        block = []
        for label, (c1, c2, _) in COG_BASE.items():
            col = c1 if phase == "baseline" else c2
            by = {g: [r.get(col) for r in subset(rows, g)] for g in (1, 2, 3)}
            ow = oneway(by)
            if ow["method"] == "insufficient n":
                p = None
            elif ow["equal_var"]:
                p = ow["P"]
            else:
                p = welch_anova(by)
            block.append(
                {
                    "outcome": label,
                    "active": fmt_ms(*mean_sd(by[1])[:2]),
                    "sham": fmt_ms(*mean_sd(by[2])[:2]),
                    "control": fmt_ms(*mean_sd(by[3])[:2]),
                    "F": round(ow["F"], 3) if ow.get("F") is not None and ow["equal_var"] else None,
                    "P": round(p, 3) if p is not None and p == p else None,
                    "method": ow["method"],
                }
            )
        result["tables"][f"table_{phase}"] = block

    # Post-hoc for post scores on TMT/VFT/HVLT
    post_pair = {}
    effect = {}
    for label in ["TMT-B", "VFT", "HVLT-R"]:
        col = COG_BASE[label][1]
        vals, labs = [], []
        for g in (1, 2, 3):
            for r in subset(rows, g):
                v = r.get(col)
                if v is not None and v != "":
                    vals.append(float(v))
                    labs.append(g)
        by = {g: [r.get(col) for r in subset(rows, g) if r.get(col) is not None] for g in (1, 2, 3)}
        ow = oneway(by)
        equal_var = ow.get("equal_var", True)
        pw = pairwise(vals, labs, equal_var)
        post_pair[label] = {
            "overall_P": round(ow["P"], 3) if ow["P"] is not None and ow["equal_var"] else round(welch_anova(by), 3),
            "active_vs_sham": round(pw[(1, 2)], 3),
            "active_vs_control": round(pw[(1, 3)], 3),
            "sham_vs_control": round(pw[(2, 3)], 3),
        }
        effect[label] = {
            "active_vs_sham_d": round(cohen_d(by[1], by[2]), 2),
            "active_vs_control_d": round(cohen_d(by[1], by[3]), 2),
        }
    result["pairwise_post"] = post_pair
    result["effect_sizes"] = effect

    post_p = {row["outcome"]: row["P"] for row in result["tables"]["table_post"] if row["P"] is not None}
    fdr = fdr_bh(post_p)
    primary = ["TMT-B", "VFT", "HVLT-R"]
    result["multiplicity"] = {
        "method": "Benjamini-Hochberg FDR across five post-intervention omnibus tests",
        "primary_outcomes": primary,
        "secondary_outcomes": ["MMSE", "DST"],
        "raw_p": post_p,
        "fdr_adjusted_p": fdr,
        "primary_significant_at_fdr_0_05": {k: fdr.get(k, 1.0) < 0.05 for k in primary},
    }

    within_group = []
    for label, (c_pre, c_post, _) in COG_BASE.items():
        for grp in (1, 2, 3):
            pre_vals, post_vals = [], []
            for r in subset(rows, grp):
                a, b = r.get(c_pre), r.get(c_post)
                if a is None or b is None or a == "" or b == "":
                    continue
                pre_vals.append(float(a))
                post_vals.append(float(b))
            if len(pre_vals) < 3:
                continue
            t_stat, p_val = stats.ttest_rel(pre_vals, post_vals)
            pre_m, pre_sd, _ = mean_sd(pre_vals)
            post_m, post_sd, _ = mean_sd(post_vals)
            within_group.append(
                {
                    "outcome": label,
                    "group": GROUP_LABEL[grp],
                    "pre": fmt_ms(pre_m, pre_sd),
                    "post": fmt_ms(post_m, post_sd),
                    "delta": round(post_m - pre_m, 2),
                    "t": round(float(t_stat), 3),
                    "P": round(float(p_val), 4) if p_val == p_val else None,
                    "n": len(pre_vals),
                }
            )
    result["within_group_change"] = within_group

    change_scores = []
    change_pairwise = {}
    for label, (c_pre, c_post, _) in COG_BASE.items():
        deltas_by = {g: [] for g in (1, 2, 3)}
        for r in rows:
            grp = r.get("分组")
            pre, post = r.get(c_pre), r.get(c_post)
            if pre is None or post is None or pre == "" or post == "":
                continue
            deltas_by[grp].append(float(post) - float(pre))
        ow = oneway(deltas_by)
        if ow.get("F") is None:
            continue
        p_val = ow["P"] if ow["equal_var"] else welch_anova(deltas_by)
        active_d = deltas_by[1]
        sham_d = deltas_by[2]
        control_d = deltas_by[3]
        row = {
            "outcome": label,
            "active_delta": round(float(np.mean(active_d)), 2) if active_d else None,
            "sham_delta": round(float(np.mean(sham_d)), 2) if sham_d else None,
            "control_delta": round(float(np.mean(control_d)), 2) if control_d else None,
            "F": round(ow["F"], 3) if ow.get("F") is not None else None,
            "P": round(p_val, 4) if p_val is not None and p_val == p_val else None,
            "method": ow["method"],
        }
        change_scores.append(row)
        if label in {"TMT-B", "VFT", "HVLT-R"}:
            vals, labs = [], []
            for g in (1, 2, 3):
                for d in deltas_by[g]:
                    vals.append(d)
                    labs.append(g)
            pw = pairwise(vals, labs, ow.get("equal_var", True))
            change_pairwise[label] = {
                "overall_P": round(p_val, 4) if p_val is not None else None,
                "active_vs_sham": round(pw[(1, 2)], 4),
                "active_vs_control": round(pw[(1, 3)], 4),
                "sham_vs_control": round(pw[(2, 3)], 4),
            }
    result["change_score_anova"] = change_scores
    result["change_score_pairwise"] = change_pairwise

    # Pearson: baseline cognition vs drinking features (all subjects)
    corr_rows = []
    drink_cols = ["首次饮酒年龄", "累计使用时间", "成瘾时间", "戒酒次数"]
    cog_cols = ["MMSE1", "TMT-B1", "DST1", "VFT1", "HVLT-R1"]
    for dc in drink_cols:
        for cc in cog_cols:
            xs, ys = [], []
            for r in rows:
                x, y = r.get(dc), r.get(cc)
                if x is None or y is None or x == "" or y == "":
                    continue
                xs.append(float(x))
                ys.append(float(y))
            if len(xs) < 5:
                continue
            r_val, p_val = stats.pearsonr(xs, ys)
            if p_val < 0.05:
                corr_rows.append(
                    {
                        "drinking": dc,
                        "cognition": cc.replace("1", ""),
                        "r": round(float(r_val), 3),
                        "P": round(float(p_val), 3),
                        "n": len(xs),
                    }
                )
    result["correlation"]["significant"] = corr_rows

    return result


def sync_group_files(headers: list[str], rows: list[dict]) -> None:
    """Rewrite group workbooks from cleaned master."""
    splits = {
        "真刺激组比较.xlsx": 1,
        "伪刺激组比较.xlsx": 2,
        "酒依赖比较.xlsx": 3,
    }
    for fname, grp in splits.items():
        path = TABLE_DIR / fname
        sub = [r for r in rows if r.get("分组") == grp]
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws.delete_rows(1, ws.max_row)
        ws.append(headers)
        for r in sub:
            ws.append([r.get(h) for h in headers])
        wb.save(path)

    # legacy single-group sheets without 分组 column
    legacy_map = {
        "真刺激组.xlsx": 1,
        "伪刺激组.xlsx": 2,
        "酒依赖组.xlsx": 3,
    }
    legacy_headers = [h for h in headers if h != "分组"]
    for fname, grp in legacy_map.items():
        path = TABLE_DIR / fname
        if not path.exists():
            continue
        sub = [r for r in rows if r.get("分组") == grp]
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws.delete_rows(1, ws.max_row)
        ws.append(legacy_headers)
        for r in sub:
            ws.append([r.get(h) for h in legacy_headers])
        wb.save(path)


def main() -> None:
    headers, rows = load_compare()
    cleaned, audit = clean_compare_rows(rows)
    save_compare(headers, cleaned)
    sync_group_files(headers, cleaned)

    stats_out = build_tables(cleaned)
    stats_out["audit"] = audit
    stats_out["total_n"] = len(cleaned)

    out_json = OUT_DIR / "analysis_results.json"
    out_json.write_text(json.dumps(stats_out, ensure_ascii=False, indent=2), encoding="utf-8")

    print("Cleaned n:", stats_out["n"], "total", len(cleaned))
    print("Audit:", json.dumps(audit, ensure_ascii=False, indent=2))
    print("Saved:", out_json)


if __name__ == "__main__":
    main()
